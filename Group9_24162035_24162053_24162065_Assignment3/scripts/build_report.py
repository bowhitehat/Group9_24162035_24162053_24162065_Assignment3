from pathlib import Path

from PIL import Image
from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION_START
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
EVIDENCE = ROOT / "output" / "qa-evidence" / "framed"
OUT = ROOT / "output" / "report" / "Bao_cao_Assignment3_Nhom9.docx"
OUT.parent.mkdir(parents=True, exist_ok=True)

BLUE = "17365D"
MID_BLUE = "2F75B5"
LIGHT_BLUE = "DDEBF7"
PALE_BLUE = "F3F8FC"
GREEN = "00875A"
GRAY = "5B6573"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_cell_text(cell, text, bold=False, color="202A35", size=8.0, align=WD_ALIGN_PARAGRAPH.LEFT):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.space_before = Pt(0)
    r = p.add_run(text)
    r.bold = bold
    r.font.name = "Arial"
    r.font.size = Pt(size)
    r.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_hyperlink(paragraph, text, url):
    part = paragraph.part
    rel_id = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rel_id)
    run = OxmlElement("w:r")
    props = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), MID_BLUE)
    props.append(color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    props.append(underline)
    run.append(props)
    text_node = OxmlElement("w:t")
    text_node.text = text
    run.append(text_node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run("Trang ")
    run.font.name = "Arial"
    run.font.size = Pt(9)
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = "PAGE"
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char1, instr_text, fld_char2])


def style_section(section, landscape=False):
    if landscape:
        section.orientation = WD_ORIENT.LANDSCAPE
        section.page_width = Inches(11)
        section.page_height = Inches(8.5)
        margin = Inches(0.45)
    else:
        section.orientation = WD_ORIENT.PORTRAIT
        section.page_width = Inches(8.5)
        section.page_height = Inches(11)
        margin = Inches(0.65)
    section.top_margin = margin
    section.bottom_margin = margin
    section.left_margin = margin
    section.right_margin = margin
    section.header_distance = Inches(0.25)
    section.footer_distance = Inches(0.25)


def add_body(doc, text, bold=False, color="202A35", align=None, before=0, after=5, size=10.5):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.12
    r = p.add_run(text)
    r.bold = bold
    r.font.name = "Arial"
    r.font.size = Pt(size)
    r.font.color.rgb = RGBColor.from_string(color)
    return p


def add_heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    p.paragraph_format.keep_with_next = True
    return p


doc = Document()
style_section(doc.sections[0])
add_page_number(doc.sections[0].footer.paragraphs[0])

styles = doc.styles
styles["Normal"].font.name = "Arial"
styles["Normal"].font.size = Pt(10.5)
styles["Title"].font.name = "Arial"
styles["Title"].font.size = Pt(25)
styles["Title"].font.bold = True
styles["Title"].font.color.rgb = RGBColor.from_string(BLUE)
styles["Heading 1"].font.name = "Arial"
styles["Heading 1"].font.size = Pt(16)
styles["Heading 1"].font.bold = True
styles["Heading 1"].font.color.rgb = RGBColor.from_string(BLUE)
styles["Heading 2"].font.name = "Arial"
styles["Heading 2"].font.size = Pt(12.5)
styles["Heading 2"].font.bold = True
styles["Heading 2"].font.color.rgb = RGBColor.from_string(MID_BLUE)

# Trang bìa
add_body(doc, "TRƯỜNG ĐẠI HỌC / KHOA CÔNG NGHỆ THÔNG TIN", bold=True, color=GRAY,
         align=WD_ALIGN_PARAGRAPH.CENTER, after=10, size=11)
add_body(doc, "NHÓM 9", bold=True, color=MID_BLUE, align=WD_ALIGN_PARAGRAPH.CENTER, after=34, size=15)

title = doc.add_paragraph(style="Title")
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
title.add_run("BÁO CÁO KIỂM THỬ\nHỆ THỐNG ĐẶT PHÒNG HỌP")
add_body(doc, "Assignment 03", bold=True, color=GRAY, align=WD_ALIGN_PARAGRAPH.CENTER, after=28, size=14)

info = doc.add_table(rows=4, cols=2)
info.alignment = WD_TABLE_ALIGNMENT.CENTER
info.autofit = False
info.columns[0].width = Inches(1.65)
info.columns[1].width = Inches(4.8)
cover_rows = [
    ("Đề tài", "Booking Management - Spring Boot / Thymeleaf"),
    ("Thành viên 1", "Vũ Trọng Hưng - MSSV 24162035"),
    ("Thành viên 2", "Trang Sĩ Hoàng - MSSV 24162053"),
    ("Thành viên 3", "MSSV 24162065"),
]
for idx, (label, value) in enumerate(cover_rows):
    set_cell_text(info.cell(idx, 0), label, bold=True, color="FFFFFF", size=10)
    set_cell_shading(info.cell(idx, 0), BLUE)
    set_cell_text(info.cell(idx, 1), value, size=10)
    set_cell_shading(info.cell(idx, 1), PALE_BLUE if idx % 2 == 0 else "FFFFFF")

add_body(doc, "", after=6)
repo_p = doc.add_paragraph()
repo_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
repo_p.paragraph_format.space_before = Pt(8)
repo_p.paragraph_format.space_after = Pt(6)
r = repo_p.add_run("GitHub repository: ")
r.bold = True
r.font.name = "Arial"
r.font.size = Pt(10)
add_hyperlink(repo_p, "Group9_24162035_24162053_24162065_Assignment3", "https://github.com/bowhitehat/Group9_24162035_24162053_24162065_Assignment3")
add_body(doc, "Ngày kiểm thử: 23/09/2026", color=GRAY, align=WD_ALIGN_PARAGRAPH.CENTER, before=10, size=10)
doc.add_page_break()

# Mục 1
add_heading(doc, "1. KIẾN TRÚC VÀ CÔNG NGHỆ", 1)
add_body(doc, "Ứng dụng được tổ chức theo luồng Controller -> Service -> DTO/Model. BookingController nhận request HTTP, thực hiện form binding và điều hướng view. BookingService tập trung toàn bộ quy tắc nghiệp vụ, còn BookingFormDto và Booking biểu diễn dữ liệu nhập và dữ liệu miền.")

flow = doc.add_table(rows=1, cols=4)
flow.alignment = WD_TABLE_ALIGNMENT.CENTER
flow.autofit = False
for i, text in enumerate(["HTTP / Thymeleaf", "BookingController", "BookingService", "ArrayList<Booking>"]):
    set_cell_text(flow.cell(0, i), text, bold=True, color="FFFFFF", size=9, align=WD_ALIGN_PARAGRAPH.CENTER)
    set_cell_shading(flow.cell(0, i), BLUE if i % 2 == 0 else MID_BLUE)

add_heading(doc, "1.1 Controller và giao diện", 2)
add_body(doc, "BookingController cung cấp các chức năng liệt kê, tạo mới, chỉnh sửa và hủy booking. Giao diện Thymeleaf liên kết dữ liệu qua BookingFormDto, hiển thị lỗi theo từng field và thông báo nghiệp vụ trên trang danh sách.")
add_heading(doc, "1.2 Service và lưu trữ in-memory", 2)
add_body(doc, "BookingService kiểm tra thời gian bắt đầu, thứ tự startAt/endAt, thời lượng tối đa 120 phút, trùng lịch cùng phòng, giới hạn tối đa 2 booking CONFIRMED cho một người và điều kiện hủy trước ít nhất 30 phút. Dữ liệu được lưu bằng ArrayList trong bộ nhớ; bản ghi hủy đổi trạng thái thành CANCELLED thay vì bị xóa.")
add_heading(doc, "1.3 Công nghệ sử dụng", 2)
add_body(doc, "Spring Boot 3.3.4, Spring MVC, Thymeleaf, Jakarta Bean Validation, Maven, Java, JUnit 5 và MockMvc. Bộ test tự động đã chạy thành công 16/16 test, gồm 9 service tests và 7 controller tests.")

# Mục 2
add_heading(doc, "2. BẢNG PHÂN CÔNG CÔNG VIỆC", 1)
assign = doc.add_table(rows=1, cols=3)
assign.style = "Table Grid"
assign.alignment = WD_TABLE_ALIGNMENT.CENTER
for cell, text in zip(assign.rows[0].cells, ["Thành viên", "MSSV", "Đóng góp"]):
    set_cell_text(cell, text, bold=True, color="FFFFFF", size=9.5, align=WD_ALIGN_PARAGRAPH.CENTER)
    set_cell_shading(cell, BLUE)
assignment_rows = [
    ("Vũ Trọng Hưng", "24162035", "Setup kiến trúc Spring Boot, Model, DTO, Service logic in-memory, Unit Tests 9/9 rules."),
    ("Trang Sĩ Hoàng", "24162053", "Xây dựng giao diện Thymeleaf, xử lý form binding và hiển thị lỗi từng field."),
    ("Thành viên 3", "24162065", "Thiết kế kịch bản test, chạy kiểm thử thực tế, chụp ảnh minh chứng và viết báo cáo."),
]
for idx, row_data in enumerate(assignment_rows):
    cells = assign.add_row().cells
    for j, value in enumerate(row_data):
        set_cell_text(cells[j], value, size=9.2)
        set_cell_shading(cells[j], PALE_BLUE if idx % 2 == 0 else "FFFFFF")

# Mục 3 matrix ở landscape
land = doc.add_section(WD_SECTION_START.NEW_PAGE)
style_section(land, landscape=True)
add_heading(doc, "3. CHI TIẾT KẾT QUẢ KIỂM THỬ", 1)
add_body(doc, "Kết quả tổng hợp: 12/12 ca kiểm thử thủ công PASS; ứng dụng không crash. Các ảnh minh chứng được chụp trực tiếp từ hệ thống chạy tại localhost:8080 và trình bày ở các trang tiếp theo.", bold=True, color=GREEN, size=10)

headers = ["STT", "Tên rule cần test", "Dữ liệu nhập (Input)", "Kết quả mong đợi", "Kết quả thực tế", "Đánh giá"]
cases = [
    ("1", "Form rỗng", "Không nhập 5 trường; Submit", "Hiện lỗi từng ô, không crash", "Hiện đủ 5 lỗi required tại từng field; form giữ nguyên", "PASS"),
    ("2", "Tạo booking hợp lệ", "E505; Le Minh QA; 25/09 10:00-11:00", "Tạo thành công", "Thông báo thành công; sinh booking #4 CONFIRMED", "PASS"),
    ("3", "Start in the past", "P101; 22/09 10:00-11:00", "Báo start ở quá khứ", "Hiện lỗi: thời gian bắt đầu không được ở quá khứ", "PASS"),
    ("4", "Invalid interval", "P102; 25/09 11:00-10:30", "Báo endAt <= startAt", "Hiện lỗi thời gian kết thúc phải sau bắt đầu", "PASS"),
    ("5", "Duration too long", "P103; 25/09 12:00-14:30", "Báo vượt 120 phút", "Hiện lỗi tối đa 120 phút; thời lượng hiện tại 150 phút", "PASS"),
    ("6", "Same-room overlap", "A101; 24/09 09:30-10:30", "Báo trùng lịch cùng phòng", "Hiện lỗi phòng A101 bị trùng khung giờ", "PASS"),
    ("7", "Different room same time", "D404; 24/09 09:00-10:00", "Khác phòng nên tạo được", "Tạo booking #5 D404 CONFIRMED thành công", "PASS"),
    ("8", "Max 2 CONFIRMED", "Nguyen Van A; Z999; 26/09 09:00-10:00", "Từ chối booking thứ 3", "Hiện lỗi người đặt đã có tối đa 2 booking CONFIRMED", "PASS"),
    ("9", "Cancel too late", "Hủy #3 khi còn 9 phút", "Từ chối hủy, báo lỗi tại list", "Thông báo lỗi đỏ; #3 vẫn CONFIRMED", "PASS"),
    ("10", "Cancel in time", "Hủy #1 khi còn >30 phút", "Đổi CANCELLED, không xóa", "#1 vẫn tồn tại và có trạng thái CANCELLED", "PASS"),
    ("11", "Reuse freed slot", "A101; đúng slot #1 đã hủy", "Tạo booking mới thành công", "Tạo booking #6 A101 CONFIRMED; #1 vẫn CANCELLED", "PASS"),
    ("12", "Edit confirmed booking", "Sửa #2: B202 -> B203", "Cập nhật hợp lệ", "Thông báo cập nhật thành công; #2 thành B203", "PASS"),
]
matrix = doc.add_table(rows=1, cols=6)
matrix.style = "Table Grid"
matrix.alignment = WD_TABLE_ALIGNMENT.CENTER
matrix.autofit = False
widths = [0.45, 1.55, 1.75, 1.9, 3.05, 0.75]
for cell, text, width in zip(matrix.rows[0].cells, headers, widths):
    cell.width = Inches(width)
    set_cell_text(cell, text, bold=True, color="FFFFFF", size=7.6, align=WD_ALIGN_PARAGRAPH.CENTER)
    set_cell_shading(cell, BLUE)
set_repeat_table_header(matrix.rows[0])
for idx, row_data in enumerate(cases):
    cells = matrix.add_row().cells
    for j, (cell, value) in enumerate(zip(cells, row_data)):
        set_cell_text(cell, value, bold=(j == 5), color=(GREEN if j == 5 else "202A35"), size=7.2,
                      align=(WD_ALIGN_PARAGRAPH.CENTER if j in (0, 5) else WD_ALIGN_PARAGRAPH.LEFT))
        set_cell_shading(cell, PALE_BLUE if idx % 2 == 0 else "FFFFFF")

# Quay lại portrait cho ảnh minh chứng
portrait = doc.add_section(WD_SECTION_START.NEW_PAGE)
style_section(portrait, landscape=False)
add_heading(doc, "3.1 ẢNH MINH CHỨNG CHI TIẾT", 1)
add_body(doc, "Mỗi ảnh thể hiện URL chạy thực tế, dữ liệu/record liên quan, thông báo lỗi hoặc thông báo thành công và trạng thái sau thao tác.", color=GRAY, size=10)

evidence_notes = [
    ("TC01 - Form rỗng", "Submit biểu mẫu không nhập dữ liệu. Năm thông báo lỗi hiển thị ngay dưới các trường; hệ thống không crash."),
    ("TC02 - Tạo booking hợp lệ", "Booking E505 được tạo thành công và xuất hiện trên danh sách với trạng thái CONFIRMED."),
    ("TC03 - Start in the past", "Ứng dụng từ chối thời gian bắt đầu 22/09/2026 vì đã nằm trong quá khứ."),
    ("TC04 - Invalid interval", "endAt 10:30 nhỏ hơn startAt 11:00; lỗi nghiệp vụ hiển thị tại form."),
    ("TC05 - Duration too long", "Khoảng thời gian 150 phút vượt giới hạn 120 phút và bị từ chối."),
    ("TC06 - Same-room overlap", "Phòng A101 trùng với booking hiện hữu trong cùng khung giờ nên không được tạo."),
    ("TC07 - Different room same time", "Phòng D404 được đặt cùng giờ với A101 vì khác phòng; booking #5 được tạo thành công."),
    ("TC08 - Max 2 CONFIRMED bookings", "Nguyen Van A đã có 2 booking CONFIRMED; hệ thống từ chối booking thứ ba."),
    ("TC09 - Cancel too late", "Yêu cầu hủy booking #3 khi chỉ còn 9 phút bị từ chối; lỗi đỏ hiển thị trên trang list."),
    ("TC10 - Cancel in time", "Booking #1 được đổi sang CANCELLED, record vẫn còn và các thao tác sửa/hủy bị khóa."),
    ("TC11 - Reuse freed slot", "Booking #6 tái sử dụng đúng phòng A101 và khung giờ của record #1 đã CANCELLED."),
    ("TC12 - Edit confirmed booking", "Booking #2 được cập nhật phòng B202 thành B203; thông báo thành công xuất hiện trên danh sách."),
]

for i, (heading, note) in enumerate(evidence_notes, start=1):
    add_heading(doc, heading, 2)
    note_p = add_body(doc, note, size=9.8, after=7)
    note_p.paragraph_format.keep_with_next = True
    matches = sorted(EVIDENCE.glob(f"{i:02d}-*.png"))
    if not matches:
        raise FileNotFoundError(f"Missing evidence image for case {i:02d}")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    with Image.open(matches[0]) as screenshot:
        px_w, px_h = screenshot.size
    max_w, max_h = 6.95, 7.55
    display_w = max_w
    display_h = display_w * px_h / px_w
    if display_h > max_h:
        display_h = max_h
        display_w = display_h * px_w / px_h
    p.add_run().add_picture(str(matches[0]), width=Inches(display_w), height=Inches(display_h))
    cap = add_body(doc, f"Hình {i}. Minh chứng {heading}", color=GRAY,
                   align=WD_ALIGN_PARAGRAPH.CENTER, after=0, size=8.5)
    cap.runs[0].italic = True
    if i < len(evidence_notes):
        doc.add_page_break()

# Phụ lục
doc.add_page_break()
add_heading(doc, "PHỤ LỤC - KẾT QUẢ TEST TỰ ĐỘNG", 1)
auto = doc.add_table(rows=4, cols=2)
auto.style = "Table Grid"
auto.alignment = WD_TABLE_ALIGNMENT.CENTER
auto_rows = [
    ("Lệnh thực thi", "mvn -Dmaven.repo.local=<temp> test"),
    ("Service tests", "9/9 PASS"),
    ("Controller tests", "7/7 PASS"),
    ("Tổng", "16 tests - 0 failures - 0 errors - 0 skipped - BUILD SUCCESS"),
]
for idx, (label, value) in enumerate(auto_rows):
    set_cell_text(auto.cell(idx, 0), label, bold=True, color="FFFFFF", size=9.5)
    set_cell_shading(auto.cell(idx, 0), BLUE)
    set_cell_text(auto.cell(idx, 1), value, bold=(idx == 3), color=(GREEN if idx == 3 else "202A35"), size=9.5)
    set_cell_shading(auto.cell(idx, 1), PALE_BLUE if idx % 2 == 0 else "FFFFFF")

add_body(doc, "Kết luận: toàn bộ rule bắt buộc đã được kiểm tra bằng test tự động và kiểm thử thao tác thực tế. Kết quả tại thời điểm chạy đều PASS.", bold=True, color=GREEN, before=10, size=10.5)

doc.core_properties.title = "Báo cáo kiểm thử Assignment 3 - Nhóm 9"
doc.core_properties.subject = "Booking Management QA Report"
doc.core_properties.author = "Nhóm 9"
doc.save(OUT)
print(OUT)
